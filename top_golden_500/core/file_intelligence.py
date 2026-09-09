from __future__ import annotations

import io
import re
from pathlib import Path
from typing import Iterable

import pandas as pd
from docx import Document
from pypdf import PdfReader

EMAIL_RE = re.compile(r"[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}", re.I)
PHONE_RE = re.compile(r"(?:\+?\d[\d\s().-]{7,}\d)")

COLUMN_ALIASES = {
    "company": ["company", "company name", "business", "organization", "organisation", "account", "client", "شركة", "اسم الشركة"],
    "country": ["country", "market", "location", "الدولة", "بلد"],
    "customer_type": ["type", "customer type", "buyer type", "segment", "business type", "نوع العميل", "القطاع"],
    "contact_name": ["contact", "contact name", "name", "decision maker", "person", "اسم", "اسم الشخص"],
    "job_title": ["title", "job title", "position", "role", "designation", "المسمى", "الوظيفة"],
    "email": ["email", "e-mail", "mail", "البريد", "الايميل", "الإيميل"],
    "phone": ["phone", "mobile", "telephone", "tel", "whatsapp", "wa", "رقم", "موبايل", "واتساب"],
    "website": ["website", "web", "url", "domain", "site", "الموقع"],
    "notes": ["notes", "note", "comments", "description", "activity", "details", "ملاحظات"],
}
TARGET_COLUMNS = list(COLUMN_ALIASES.keys())


def _norm(value) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip()


def _normalize_columns(df: pd.DataFrame) -> pd.DataFrame:
    lowered = {str(c).strip().lower(): c for c in df.columns}
    rename = {}
    for target, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            if alias.lower() in lowered:
                rename[lowered[alias.lower()]] = target
                break
    df = df.rename(columns=rename).copy()
    for c in TARGET_COLUMNS:
        if c not in df.columns:
            df[c] = ""
    return df[TARGET_COLUMNS]


def _rows_from_text(text: str, source_name: str) -> pd.DataFrame:
    lines = [_norm(x) for x in text.splitlines() if _norm(x)]
    rows, seen = [], set()

    for i, line in enumerate(lines):
        emails = EMAIL_RE.findall(line)
        if not emails:
            continue
        phones = PHONE_RE.findall(line)
        context = " | ".join(lines[max(0, i - 2): min(len(lines), i + 3)])
        for email in emails:
            email = email.lower().strip()
            if email in seen:
                continue
            seen.add(email)
            domain = email.split("@")[-1]
            before = line.split(email)[0].strip(" -|,;:")
            company = before[:100] or domain.split(".")[0].replace("-", " ").title()
            rows.append({
                "company": company,
                "country": "",
                "customer_type": "",
                "contact_name": "",
                "job_title": "",
                "email": email,
                "phone": phones[0] if phones else "",
                "website": f"https://{domain}",
                "notes": f"{context[:700]} | source={source_name}",
            })

    if not rows:
        for i, line in enumerate(lines):
            phones = PHONE_RE.findall(line)
            if phones:
                context = " | ".join(lines[max(0, i - 1): min(len(lines), i + 2)])
                rows.append({
                    "company": line[:100],
                    "country": "",
                    "customer_type": "",
                    "contact_name": "",
                    "job_title": "",
                    "email": "",
                    "phone": phones[0],
                    "website": "",
                    "notes": f"{context[:700]} | source={source_name}",
                })

    return pd.DataFrame(rows, columns=TARGET_COLUMNS)


def parse_upload(uploaded) -> pd.DataFrame:
    suffix = Path(uploaded.name).suffix.lower()
    frames = []

    if suffix in {".xlsx", ".xlsm", ".xls"}:
        raw = uploaded.getvalue()
        xls = pd.ExcelFile(io.BytesIO(raw))
        for sheet in xls.sheet_names:
            df = pd.read_excel(io.BytesIO(raw), sheet_name=sheet)
            if not df.empty:
                frames.append(df)
    elif suffix == ".csv":
        raw = uploaded.getvalue()
        parsed = None
        for enc in ("utf-8-sig", "utf-8", "cp1252", "latin1"):
            try:
                parsed = pd.read_csv(io.BytesIO(raw), encoding=enc)
                break
            except Exception:
                continue
        if parsed is not None:
            frames.append(parsed)
    elif suffix == ".docx":
        doc = Document(io.BytesIO(uploaded.getvalue()))
        text_parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                text_parts.append(" | ".join(cell.text for cell in row.cells))
        frames = [_rows_from_text("\n".join(text_parts), uploaded.name)]
    elif suffix == ".pdf":
        reader = PdfReader(io.BytesIO(uploaded.getvalue()))
        text = "\n".join((page.extract_text() or "") for page in reader.pages)
        frames = [_rows_from_text(text, uploaded.name)]
    else:
        raise ValueError(f"Unsupported file type: {suffix}")

    normalized = []
    for df in frames:
        if df is None or df.empty:
            continue
        normalized.append(_normalize_columns(df))
    if not normalized:
        return pd.DataFrame(columns=TARGET_COLUMNS)

    out = pd.concat(normalized, ignore_index=True)
    for c in TARGET_COLUMNS:
        out[c] = out[c].fillna("").map(_norm)
    return out


def parse_many(files: Iterable) -> pd.DataFrame:
    frames = []
    for f in files:
        df = parse_upload(f)
        if not df.empty:
            df["source_file"] = f.name
            frames.append(df)
    if not frames:
        return pd.DataFrame(columns=TARGET_COLUMNS + ["source_file"])

    out = pd.concat(frames, ignore_index=True)
    out["_email_key"] = out["email"].str.lower().str.strip()
    out["_fallback"] = (
        out["company"].str.lower().str.strip()
        + "|"
        + out["phone"].str.replace(r"\D", "", regex=True)
    )
    out["_key"] = out["_email_key"].where(out["_email_key"].ne(""), out["_fallback"])
    out = out.drop_duplicates("_key", keep="first").drop(columns=["_email_key", "_fallback", "_key"])
    return out.reset_index(drop=True)
